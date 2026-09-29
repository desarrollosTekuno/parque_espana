from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

path = 'database/seeders/data/Plantilla_Migracion_Unificada.xlsx'
book = load_workbook(path)

opening = book['05 SALDOS INICIALES']
opening.title = '05 SALDOS AL INICIO'
opening['A1'] = '1. SALDOS AL INICIO DEL AÑO   (billing.charges / billing.credit_balances)'
opening['A3'] = 'Fecha de corte anterior al primer cargo del año que se va a migrar; por ejemplo, 2025-12-31 para pagos desde 2026-01-01.'
opening['D3'] = 'Deuda anterior al inicio del año. No incluya cargos generados desde enero.'
opening['E3'] = 'De esa deuda anterior, cuanto ya estaba vencido al corte.'
opening['F3'] = 'Saldo a favor existente antes del inicio del año.'
opening['H2'] = 'REFERENCIA DEL SALDO INICIAL'
opening['H3'] = 'Clave sencilla para relacionar un pago posterior con este saldo, por ejemplo INI-1001-PE1. Una por cuenta y club.'
opening['H2']._style = copy(opening['B2']._style)
opening['H3']._style = copy(opening['B3']._style)
opening.column_dimensions['H'].width = 30
opening.unmerge_cells('A1:G1')
opening.merge_cells('A1:H1')

old = book['05 CARGOS PENDIENTES']
old.title = '05 CARGOS ANTERIORES'
old['A1'] = '2. CARGOS ANTERIORES SIN PAGAR   (billing.charges)'
old.unmerge_cells('A1:N1')
old.delete_cols(4)
old.merge_cells('A1:M1')
old['A3'] = 'Folio del cargo anterior al inicio del año. Si no existe, asigne una referencia unica.'
old['F3'] = 'Importe original del cargo anterior, en MXN.'
old['G3'] = 'Saldo que seguia pendiente antes del inicio del año, en MXN.'
old['M3'] = 'PENDIENTE o PARCIAL al inicio del año. No incluya cargos ya pagados.'

instructions = book['INSTRUCCIONES']
instructions['B7'] = ('Para conservar pagos desde enero: capture la deuda previa en SALDOS AL INICIO o en CARGOS ANTERIORES, '
                      'luego los CARGOS DEL AÑO, PAGOS DEL AÑO y APLICACIONES. No repita el mismo adeudo en ambas hojas de apertura.')
instructions['A12'] = 'PAGOS HISTORICOS'
instructions['B12'] = ('Un pago puede cubrir varios cargos y un cargo puede recibir varios pagos. '
                       'En APLICACIONES registre una fila por cada relacion. Un pago sin aplicacion queda sin asignar a un cargo.')
instructions['A12'].font = copy(instructions['A7'].font)
instructions['A12'].fill = copy(instructions['A7'].fill)


def add_sheet(name, title, columns):
    sheet = book.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A4'
    sheet['A1'] = title
    sheet['A1'].fill = PatternFill('solid', fgColor='FF1A252F')
    sheet['A1'].font = Font(name='Arial', size=12, bold=True, color='FFFFFFFF')
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 28
    sheet.row_dimensions[3].height = 37
    for col, (header, hint, required) in enumerate(columns, 1):
        letter = get_column_letter(col)
        a = sheet.cell(2, col, header)
        a.fill = PatternFill('solid', fgColor='FFC0392B' if required else 'FF2980B9')
        a.font = Font(name='Arial', size=10, bold=True, color='FFFFFFFF')
        a.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        b = sheet.cell(3, col, hint)
        b.fill = PatternFill('solid', fgColor='FFECF0F1')
        b.font = Font(name='Arial', size=8, italic=True, color='FF555555')
        b.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.column_dimensions[letter].width = max(20, min(36, len(header) + 6))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    return sheet


charges = add_sheet('05 CARGOS DEL AÑO', '3. CARGOS DEL AÑO   (billing.charges)', [
    ('FOLIO DEL CARGO', 'Folio del sistema anterior. Si no existe, asigne uno como CAR-0001.', True),
    ('NÚMERO DE MEMBRESÍA', 'Debe existir en 04 CUENTAS.', True),
    ('NOMBRE DEL CLUB', 'Ver 02 CLUBES.', True),
    ('CONCEPTO DE COBRO', 'Use el nombre que aparece en 01 CATALOGOS.', True),
    ('DESCRIPCIÓN', 'Detalle corto del cargo, si hace falta.', False),
    ('IMPORTE ORIGINAL', 'Monto total del cargo en MXN antes de pagos.', True),
    ('FECHA DE EMISIÓN', 'Formato AAAA-MM-DD. Desde el primer dia del año a migrar.', True),
    ('FECHA DE VENCIMIENTO', 'Formato AAAA-MM-DD, si se conoce.', False),
    ('AÑO DEL PERIODO', 'Año de la cuota, si aplica.', False),
    ('MES DEL PERIODO', 'Mes de la cuota del 1 al 12, si aplica.', False),
    ('ESTATUS', 'PENDIENTE, PARCIAL, PAGADO o CANCELADO.', True),
])

payments = add_sheet('05 PAGOS DEL AÑO', '4. PAGOS DEL AÑO   (billing.payments)', [
    ('FOLIO DEL PAGO', 'Folio del sistema anterior. Si no existe, asigne uno como PAG-0001.', True),
    ('NÚMERO DE MEMBRESÍA', 'Debe existir en 04 CUENTAS.', True),
    ('NOMBRE DEL CLUB', 'Club que recibio el pago. Ver 02 CLUBES.', True),
    ('FECHA Y HORA DEL PAGO', 'Formato AAAA-MM-DD HH:MM. Desde el primer dia del año a migrar.', True),
    ('MÉTODO DE PAGO', 'Use el nombre que aparece en 01 CATALOGOS.', True),
    ('IMPORTE', 'Monto total recibido en MXN.', True),
    ('REFERENCIA BANCARIA', 'Llenar solo si el metodo la tiene.', False),
    ('BANCO', 'Llenar solo si aplica.', False),
    ('NÚMERO DE CHEQUE', 'Llenar solo para cheque.', False),
    ('ESTATUS', 'REGISTRADO o CANCELADO.', True),
    ('NOTAS', 'Comentario opcional sobre el pago.', False),
])

applications = add_sheet('05 APLICACIONES DE PAGOS', '5. APLICACIONES DE PAGOS   (billing.payment_applications)', [
    ('FOLIO DEL PAGO', 'Debe coincidir con 05 PAGOS DEL AÑO.', True),
    ('FOLIO DEL CARGO', 'Debe coincidir con 05 CARGOS DEL AÑO, 05 CARGOS ANTERIORES o la referencia del saldo inicial.', True),
    ('IMPORTE APLICADO', 'Parte del pago destinada a este cargo, en MXN. Use varias filas si cubre varios cargos.', True),
    ('DESCUENTO', 'Descuento aplicado a este cargo, si existio. En MXN.', False),
])

# Keep the capture steps together and before collection notes.
for sheet in (charges, payments, applications):
    book._sheets.remove(sheet)
position = book.sheetnames.index('05 NOTAS DE COBRANZA')
for offset, sheet in enumerate((charges, payments, applications)):
    book._sheets.insert(position + offset, sheet)

examples = book['EJEMPLOS']
examples.unmerge_cells('A127:G127')
examples.merge_cells('A127:H127')
examples['A127'] = 'EJEMPLOS: 05 SALDOS AL INICIO'
examples['A133'] = 'EJEMPLOS: 05 CARGOS ANTERIORES'
for row in (128, 134):
    target = opening if row == 128 else old
    for col in range(1, target.max_column + 1):
        examples.cell(row, col, target.cell(2, col).value)
        examples.cell(row, col)._style = copy(target.cell(2, col)._style)
for row in (129, 130, 135, 136):
    for col in range(1, 15):
        examples.cell(row, col).value = None
examples['A129'] = '2025-12-31'
examples['B129'] = 'M-1001'
examples['C129'] = 'Parque España I'
examples['D129'] = 300
examples['E129'] = 300
examples['F129'] = 0
examples['H129'] = 'INI-1001-PE1'
for col, value in enumerate(['ANT-001','M-1002','Parque España I','CUOTA MENSUALIDAD','Cuota diciembre anterior',850,850,'2025-12-01','2025-12-31',2025,12,'SÍ','PENDIENTE'],1):
    examples.cell(135,col,value)

next_row = examples.max_row + 2
for sheet, sample in [
    (charges, ['CAR-001','M-1001','Parque España I','CUOTA MENSUALIDAD','Cuota enero',850,'2026-01-01','2026-01-31',2026,1,'PAGADO']),
    (payments, ['PAG-001','M-1001','Parque España I','2026-01-10 10:00','Efectivo',850,None,None,None,'REGISTRADO',None]),
    (applications, ['PAG-001','CAR-001',850,None]),
]:
    examples.cell(next_row, 1, f'EJEMPLOS: {sheet.title}')
    examples.cell(next_row, 1).fill = PatternFill('solid', fgColor='FF1A252F')
    examples.cell(next_row, 1).font = Font(name='Arial', bold=True, color='FFFFFFFF')
    examples.merge_cells(start_row=next_row, start_column=1, end_row=next_row, end_column=sheet.max_column)
    for col in range(1, sheet.max_column + 1):
        a = examples.cell(next_row + 1, col, sheet.cell(2, col).value)
        a._style = copy(sheet.cell(2, col)._style)
        examples.cell(next_row + 2, col, sample[col - 1])
    next_row += 5

book.save(path)
print('updated', path, 'tabs', len(book.sheetnames))
