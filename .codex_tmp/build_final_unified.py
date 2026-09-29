import gc
import glob
import os
import re
import unicodedata
from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = 'database/seeders/data'
OUT = os.path.join(ROOT, 'Plantilla_Migracion_Unificada.xlsx')
FILES = sorted(glob.glob(os.path.join(ROOT, '0*.xlsx')))


def wording(value):
    if not isinstance(value, str):
        return value
    references = {
        '01_CATALOGOS': '01 CATALOGOS',
        '02_CLUBES_Y_CONFIGURACION': '02 CLUBES',
        '03_SOCIOS': '03 USUARIOS',
        '04_CUENTAS_Y_MEMBRESIAS': '04 CUENTAS',
        '05_FINANZAS': '05 SALDOS INICIALES',
        '06_CASILLEROS': '06 CASILLEROS',
        '07_AMENIDADES_Y_CLASES': '07 AMENIDADES',
        '08_RESERVACIONES_Y_ACCESOS': '08 RESERVACIONES',
        '09_INFORMACION_ADMINISTRATIVA': '09 ACTAS',
    }
    for old, new in references.items():
        value = value.replace(old, new)
    def replace(match):
        word = 'usuarios' if match.group().lower() == 'socios' else 'usuario'
        return word.upper() if match.group().isupper() else word.capitalize() if match.group().istitle() else word
    return re.sub(r'\bsocios?\b', replace, value, flags=re.I)


def copy_style(source, target):
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format


def relevant_columns(ws, is_main_user):
    headers = [ws.cell(2, col).value for col in range(1, ws.max_column + 1)]
    has_number = any(isinstance(h, str) and 'NÚMERO DE SOCIO' in h for h in headers)
    keep = []
    for i, h in enumerate(headers, 1):
        if h is None:
            continue
        if has_number and not is_main_user and h in ('NOMBRE', 'APELLIDO PATERNO', 'APELLIDO MATERNO'):
            continue
        keep.append(i)
    return keep


def copy_capture(src, dest, is_main_user=False):
    keep = relevant_columns(src, is_main_user)
    for new_col, old_col in enumerate(keep, 1):
        column = get_column_letter(new_col)
        old_column = get_column_letter(old_col)
        dest.column_dimensions[column].width = src.column_dimensions[old_column].width or 18
        for row in range(1, src.max_row + 1):
            old = src.cell(row, old_col)
            new = dest.cell(row, new_col)
            new.value = wording(old.value) if row <= 3 else old.value
            copy_style(old, new)
    for row in (1, 2, 3):
        dest.row_dimensions[row].height = src.row_dimensions[row].height
    if keep:
        dest.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(keep))
    if is_main_user:
        dest['A1'] = '1. USUARIOS   (members.members)'
    dest.freeze_panes = 'A4'
    dest.sheet_view.showGridLines = src.sheet_view.showGridLines
    return keep


def add_catalog_block(src, dest, start_row):
    keep = relevant_columns(src, False)
    for new_col, old_col in enumerate(keep, 1):
        old_letter = get_column_letter(old_col)
        new_letter = get_column_letter(new_col)
        dest.column_dimensions[new_letter].width = max(dest.column_dimensions[new_letter].width or 0, src.column_dimensions[old_letter].width or 18)
        for old_row in range(1, src.max_row + 1):
            old = src.cell(old_row, old_col)
            new = dest.cell(start_row + old_row - 1, new_col)
            new.value = wording(old.value) if old_row <= 3 else old.value
            copy_style(old, new)
    for old_row in (1, 2, 3):
        dest.row_dimensions[start_row + old_row - 1].height = src.row_dimensions[old_row].height
    if keep:
        dest.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=len(keep))
    return start_row + max(src.max_row, 3) + 2


def sheet_name(section, title, used):
    title = 'USUARIOS' if section == '03' and title == 'SOCIOS' else title
    title = wording(title)
    base = f'{section} {title}'[:31]
    name = base
    suffix = 2
    while name in used:
        name = f'{base[:28]} {suffix}'
        suffix += 1
    used.add(name)
    return name


def simple_name(value):
    value = unicodedata.normalize('NFKD', str(value or ''))
    return ''.join(c for c in value.upper() if c.isalnum())


def examples_for(workbook, title, target_headers):
    if 'EJEMPLOS' not in workbook:
        return []
    ws = workbook['EJEMPLOS']
    sought = simple_name(title)
    for row in range(1, ws.max_row + 1):
        heading = ws.cell(row, 1).value
        if not isinstance(heading, str) or not heading.startswith('EJEMPLOS:'):
            continue
        block = simple_name(heading.split(':', 1)[1])
        if not (block.startswith(sought) or sought.startswith(block)):
            continue
        header_values = [wording(ws.cell(row + 1, col).value) for col in range(1, ws.max_column + 1)]
        lookup = {simple_name(h): col for col, h in enumerate(header_values, 1) if h}
        result = []
        for data_row in range(row + 2, min(row + 5, ws.max_row + 1)):
            if isinstance(ws.cell(data_row, 1).value, str) and ws.cell(data_row, 1).value.startswith('EJEMPLOS:'):
                break
            values = []
            for header in target_headers:
                col = lookup.get(simple_name(header))
                value = ws.cell(data_row, col).value if col else None
                if isinstance(value, str) and value.startswith('DATOS DE EJEMPLO'):
                    value = None
                if value == 'PARQUE ESPAÑA':
                    value = 'Parque España I'
                if value == 'PARQUE ESPAÑA 2':
                    value = 'Parque España II'
                if header in ('ACTIVA', 'ACTIVO') and isinstance(value, (int, float)):
                    value = None
                values.append(value)
            if any(v is not None for v in values):
                result.append(values)
            if len(result) == 2:
                break
        return result
    return []


wb = Workbook()
instructions = wb.active
instructions.title = 'INSTRUCCIONES'
instructions.sheet_view.showGridLines = False
instructions.merge_cells('A1:C1')
instructions['A1'] = 'INSTRUCCIONES DE CAPTURA'
instructions['A1'].fill = PatternFill('solid', fgColor='FF1A252F')
instructions['A1'].font = Font(name='Arial', size=12, bold=True, color='FFFFFFFF')
instructions.row_dimensions[1].height = 28
notes = [
    ('ORDEN', 'Siga las pestañas de izquierda a derecha: catalogos, clubes, usuarios, cuentas, finanzas e instalaciones.'),
    ('USUARIOS', 'Capture nombre y apellidos solo en la hoja USUARIOS. En las hojas relacionadas use NUMERO DE USUARIO.'),
    ('CUENTAS', 'Use NUMERO DE MEMBRESIA para relacionar cuenta, integrantes y cobros. Repita el club cuando sea necesario distinguir una membresia.'),
    ('CASILLEROS', 'Para identificar un casillero use club, numero y categoria; el mismo numero puede existir en varias categorias.'),
    ('FINANZAS', 'SALDOS INICIALES es la opcion sencilla. CARGOS PENDIENTES es el detalle alternativo. No capture el mismo adeudo en ambas hojas.'),
    ('FECHAS', 'Capture fechas como AAAA-MM-DD y horas como HH:MM.'),
    ('DATOS DESCONOCIDOS', 'Deje la celda vacia. No escriba N/A ni invente folios.'),
    ('COLORES', 'Rojo = requerido; azul = opcional. Los catalogos contienen valores de referencia y datos precargados que debe revisar.'),
    ('EJEMPLOS', 'La hoja EJEMPLOS es una guia. No copie datos ficticios a las hojas de captura.'),
]
for r, (a, b) in enumerate(notes, 3):
    instructions.cell(r, 1, a)
    instructions.cell(r, 2, b)
    instructions.cell(r, 1).font = Font(name='Arial', bold=True, color='FF1A252F')
    instructions.cell(r, 1).fill = PatternFill('solid', fgColor='FFECF0F1')
instructions.column_dimensions['A'].width = 26
instructions.column_dimensions['B'].width = 110
instructions.freeze_panes = 'A3'

catalog = wb.create_sheet('01 CATALOGOS')
catalog.freeze_panes = 'A1'
catalog.sheet_view.showGridLines = False

used = {'INSTRUCCIONES', '01 CATALOGOS', 'EJEMPLOS'}
capture_info = []
catalog_sections = []
catalog_rows = 1

# The catalog with 37,536 cities comes last so all short catalogs stay near the top.
cat_wb = load_workbook(FILES[0], rich_text=True)
catalog_source = [s for s in cat_wb if s.title != 'INSTRUCCIONES' and s.title not in ('CIUDADES', 'IMPORTES POR CLUB')]
for src in catalog_source:
    catalog_sections.append((src.title, catalog_rows))
    catalog_rows = add_catalog_block(src, catalog, catalog_rows)
cat_wb.close()
del cat_wb
gc.collect()

for path in FILES[1:]:
    section = os.path.basename(path)[:2]
    source_wb = load_workbook(path, rich_text=True)
    for src in source_wb:
        if src.title in ('INSTRUCCIONES', 'EJEMPLOS'):
            continue
        if section == '02' and src.title == 'CONCEPTOS E IMPORTES':
            pass  # Single capture place for club-specific amounts.
        if section == '04' and src.title == 'REGLAS DE PAQUETE ENTRE CLUBES':
            catalog_sections.append((src.title, catalog_rows))
            catalog_rows = add_catalog_block(src, catalog, catalog_rows)
            continue
        if section == '05' and src.title == 'MULTAS PENDIENTES':
            continue  # Captured through ACTAS and MULTAS in section 09.
        dest = wb.create_sheet(sheet_name(section, src.title, used))
        keep = copy_capture(src, dest, section == '03' and src.title == 'SOCIOS')
        if section == '03':
            dest['A1'] = wording(str(src['A1'].value))
        target_headers = [dest.cell(2, col).value for col in range(1, len(keep) + 1)]
        capture_info.append((path, src.title, dest.title, keep, examples_for(source_wb, src.title, target_headers)))
    source_wb.close()
    del source_wb
    gc.collect()

cat_wb = load_workbook(FILES[0], rich_text=True)
catalog_sections.append(('CIUDADES', catalog_rows))
catalog_rows = add_catalog_block(cat_wb['CIUDADES'], catalog, catalog_rows)
cat_wb.close()
del cat_wb
gc.collect()

# Place club rules next to club setup, without changing their source content.
rules = wb['09 REGLAS DEL CLUB']
wb._sheets.remove(rules)
club_last = max(i for i, s in enumerate(wb._sheets) if s.title.startswith('02 '))
wb._sheets.insert(club_last + 1, rules)
rules.title = '02 REGLAS DEL CLUB'

# Put the single examples sheet at the end. Each block mirrors its capture headers.
examples = wb.create_sheet('EJEMPLOS')
examples.sheet_view.showGridLines = False
examples.freeze_panes = 'A1'
example_row = 1
for path, original_name, target_name, keep, example_values in capture_info:
    target = wb[target_name if target_name in wb else '02 REGLAS DEL CLUB']
    n = len(keep)
    examples.cell(example_row, 1, f'EJEMPLOS: {target.title}')
    examples.cell(example_row, 1).fill = PatternFill('solid', fgColor='FF1A252F')
    examples.cell(example_row, 1).font = Font(name='Arial', bold=True, color='FFFFFFFF')
    examples.merge_cells(start_row=example_row, start_column=1, end_row=example_row, end_column=n)
    for col in range(1, n + 1):
        source_cell = target.cell(2, col)
        dest_cell = examples.cell(example_row + 1, col, source_cell.value)
        copy_style(source_cell, dest_cell)
        examples.column_dimensions[get_column_letter(col)].width = max(examples.column_dimensions[get_column_letter(col)].width or 0, target.column_dimensions[get_column_letter(col)].width or 18)
    if target.max_row > 3:
        example_values = [[target.cell(4 + offset, col).value for col in range(1, n + 1)] for offset in range(min(2, target.max_row - 3))]
    for offset, row_values in enumerate(example_values):
        for col, value in enumerate(row_values, 1):
            examples.cell(example_row + 2 + offset, col, value)
    example_row += 6

wb.save(OUT)
print('saved', OUT, 'sheets', len(wb.sheetnames), 'catalog_sections', len(catalog_sections), 'capture_sheets', len(capture_info))
