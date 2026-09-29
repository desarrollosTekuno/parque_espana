from copy import copy
from pathlib import Path
import re

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


ROOT = Path(r'C:\Apache24\htdocs\ParquesEsp')
catalog_path = ROOT / 'database/seeders/data/01_CATALOGOS.xlsx'
club_path = ROOT / 'database/seeders/data/02_CLUBES_Y_CONFIGURACION.xlsx'
output_path = ROOT / 'Plantilla_Migracion_Cliente.xlsx'

catalogs = openpyxl.load_workbook(catalog_path, read_only=False, data_only=False)
clubs = openpyxl.load_workbook(club_path, read_only=False, data_only=False)

order = [
    'TIPOS MEMBRESIA', 'METODOS DE PAGO', 'CONCEPTOS DE COBRO',
    'PARENTESCOS', 'ESTADOS CIVILES', 'TIPOS DOCUMENTO',
    'DOCS POR PARENTESCO', 'DOCS POR MEMBRESIA',
    'MOTIVOS CANCELACION', 'MOTIVOS SEPARACION',
    'REGLAS DE DESCUENTO', 'REGLAS DE PRECIOS',
    'ESPECIALIDADES', 'ESTATUS RESERVACION',
    'CATEGORIAS CASILLERO', 'PAISES', 'ESTADOS', 'NACIONALIDADES',
]


def no_accents(value):
    if value is None:
        return None
    text = str(value)
    replacements = str.maketrans(
        'áàäâÁÀÄÂéèëêÉÈËÊíìïîÍÌÏÎóòöôÓÒÖÔúùüûÚÙÜÛ',
        'aaaaAAAAeeeeEEEEiiiiIIIIooooOOOOuuuuUUUU',
    )
    text = text.translate(replacements)
    text = re.sub(r'01_CATALOGOS\s*-\s*', 'esta pestaña: ', text)
    text = text.replace('la hoja de Socios', 'Usuarios')
    text = text.replace('la hoja de clubes', 'la seccion CLUBES de esta pestaña')
    return text


def rows_with_values(source):
    result = []
    for source_row in source.iter_rows(min_row=4):
        if all(cell.value is None for cell in source_row):
            continue
        first = source_row[0].value
        if isinstance(first, str) and first.lstrip().startswith(('⚠', '■')):
            continue
        result.append(source_row)
    return result


book = openpyxl.Workbook()
sheet = book.active
sheet.title = 'Catalogos'
sheet.sheet_view.showGridLines = False
sheet.sheet_properties.tabColor = 'C0392B'
sheet.freeze_panes = 'A4'


def copy_style(source_cell, target_cell):
    if not source_cell.has_style:
        return
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)
    target_cell.number_format = source_cell.number_format


def place(source, target_row, target_col, name):
    data_rows = rows_with_values(source)
    width = source.max_column
    if source.title == 'CLUBES':
        width = 3

    for offset in range(width):
        title = sheet.cell(target_row, target_col + offset)
        copy_style(source.cell(1, 1), title)
        title.font = Font(name='Arial', size=10, bold=True, color='FFFFFFFF')
        if offset == 0:
            title.value = name

        for source_r, target_r in ((2, target_row + 1), (3, target_row + 2)):
            original = source.cell(source_r, offset + 1)
            target = sheet.cell(target_r, target_col + offset)
            copy_style(original, target)
            target.value = no_accents(original.value) if source_r == 3 else original.value
            if source_r == 3:
                old = target.alignment
                target.alignment = Alignment(horizontal=old.horizontal, vertical=old.vertical, wrap_text=True)

    for row_offset, source_row in enumerate(data_rows, 3):
        for col_offset in range(width):
            source_cell = source_row[col_offset]
            target = sheet.cell(target_row + row_offset, target_col + col_offset)
            target.value = source_cell.value
            copy_style(source_cell, target)

    sheet.row_dimensions[target_row].height = 21
    sheet.row_dimensions[target_row + 1].height = 29
    sheet.row_dimensions[target_row + 2].height = 55
    return target_row + 3 + len(data_rows), len(data_rows)


counts = {}
row = 1
club_sheet = clubs['CLUBES']
row, counts['CLUBES'] = place(club_sheet, row, 1, 'CLUBES')
row += 2
for name in order:
    row, counts[name] = place(catalogs[name], row, 1, name)
    row += 2

_, counts['CIUDADES'] = place(catalogs['CIUDADES'], 1, 13, 'CIUDADES')

for col in range(1, 12):
    letter = get_column_letter(col)
    original_widths = [s.column_dimensions[letter].width or 13 for s in [club_sheet] + [catalogs[n] for n in order]]
    sheet.column_dimensions[letter].width = max(18, min(40, max(original_widths)))
sheet.column_dimensions['L'].width = 3
sheet.column_dimensions['M'].width = 22
sheet.column_dimensions['N'].width = 24
sheet.column_dimensions['O'].width = 43

book.save(output_path)
print('Output:', output_path)
print('Sections:', len(counts), 'Cities:', counts['CIUDADES'])
print('Empty source sections omitted: IMPORTES POR CLUB, INCIDENTES Y SANCIONES')
